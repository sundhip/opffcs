import io
import zipfile
import pytest
import openpyxl
from fastapi.testclient import TestClient

from backend.app import app
from backend.services.exporter import TimetableExporter

@pytest.fixture
def sample_timetable():
    return {
        "timetable_id": "T_TEST_1",
        "offering_ids": ["off_cse1001_th", "off_cse1001_lab"],
        "offering_keys": ["CSE1001|DR SMITH|A1+TA1||", "CSE1001|DR SMITH||L37+L38|"],
        "per_day_timeline": {
            "MON": [
                {
                    "day": "MON",
                    "start_time": "08:00",
                    "end_time": "08:50",
                    "start_min": 480,
                    "end_min": 530,
                    "kind": "theory",
                    "slot_label": "A1",
                    "course_code": "CSE1001",
                    "faculty_name": "Dr. Smith",
                    "venue": "AB1-201"
                }
            ],
            "TUE": [
                {
                    "day": "TUE",
                    "start_time": "14:00",
                    "end_time": "15:40",
                    "start_min": 840,
                    "end_min": 940,
                    "kind": "lab",
                    "slot_label": "L37+L38",
                    "course_code": "CSE1001",
                    "faculty_name": "Dr. Smith",
                    "venue": "SJ-LAB-1"
                }
            ],
            "WED": [
                {
                    "day": "WED",
                    "start_time": "08:55",
                    "end_time": "09:45",
                    "start_min": 535,
                    "end_min": 585,
                    "kind": "theory",
                    "slot_label": "TA1",
                    "course_code": "CSE1001",
                    "faculty_name": "Dr. Smith",
                    "venue": "AB1-201"
                }
            ],
            "THU": [],
            "FRI": []
        },
        "working_days": ["MON", "TUE", "WED"],
        "free_days": ["THU", "FRI"],
        "gap_minutes_per_day": {"MON": 0, "TUE": 0, "WED": 0, "THU": 0, "FRI": 0},
        "longest_gap_minutes": 0,
        "max_daily_hours": 1.67,
        "earliest_start": "08:00",
        "latest_end": "15:40",
        "avg_faculty_rating": 4.5,
        "unrated_faculty_count": 0,
        "score": 92.5,
        "score_breakdown": {
            "gap_penalty": 0.0,
            "free_day_bonus": 20.0,
            "compactness_score": 30.0,
            "faculty_rating_score": 42.5
        }
    }

@pytest.fixture
def sample_courses():
    return [
        {
            "course_code": "CSE1001",
            "course_title": "Problem Solving and Programming",
            "credits": 4,
            "faculty_name": "Dr. Smith",
            "theory_slots": ["A1", "TA1"],
            "lab_slots": ["L37", "L38"],
            "venue": "AB1-201",
            "rating_mean": 4.5,
            "review_count": 12
        }
    ]

def test_exporter_excel_generation(sample_timetable, sample_courses):
    excel_bytes = TimetableExporter.export_to_excel(sample_timetable, sample_courses)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames
    assert "Weekly Timetable" in sheet_names
    assert "Course List" in sheet_names
    assert "Schedule Summary" in sheet_names

    ws_grid = wb["Weekly Timetable"]
    assert ws_grid.sheet_properties.pageSetUpPr.fitToPage is True
    # Merged lab cell verification: Tuesday afternoon (Row 5, columns for P7 & P8)
    merged_ranges = [str(r) for r in ws_grid.merged_cells.ranges]
    # Should include Tuesday P7-P8 merge
    assert any("5" in r for r in merged_ranges)

    ws_courses = wb["Course List"]
    assert ws_courses.cell(row=3, column=1).value == "CSE1001"
    assert ws_courses.cell(row=3, column=4).value == "Dr. Smith"

    ws_summary = wb["Schedule Summary"]
    assert ws_summary.cell(row=3, column=1).value == "Timetable ID"
    assert ws_summary.cell(row=3, column=2).value == "T_TEST_1"

def test_exporter_csv_generation(sample_timetable, sample_courses):
    csv_str = TimetableExporter.export_to_csv(sample_timetable, sample_courses)
    assert csv_str.startswith("\ufeff")  # UTF-8 BOM
    assert "CSE1001" in csv_str
    assert "Dr. Smith" in csv_str
    assert "L37+L38" in csv_str

def test_exporter_pdf_generation(sample_timetable, sample_courses):
    pdf_bytes = TimetableExporter.export_to_pdf(sample_timetable, sample_courses)
    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000

def test_exporter_png_generation(sample_timetable, sample_courses):
    png_bytes = TimetableExporter.export_to_png(sample_timetable, sample_courses, scale=2)
    assert png_bytes.startswith(b"\x89PNG")
    assert len(png_bytes) > 5000

def test_exporter_zip_bundle(sample_timetable, sample_courses):
    zip_bytes = TimetableExporter.export_zip_bundle(sample_timetable, sample_courses)
    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    namelist = zf.namelist()
    assert any(n.endswith(".xlsx") for n in namelist)
    assert any(n.endswith(".csv") for n in namelist)
    assert any(n.endswith(".pdf") for n in namelist)
    assert any(n.endswith(".png") for n in namelist)

def test_api_export_endpoints(sample_timetable, sample_courses):
    client = TestClient(app)
    payload = {"timetable": sample_timetable, "courses": sample_courses}

    # Excel
    res = client.post("/api/exports/excel", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # CSV
    res = client.post("/api/exports/csv", json=payload)
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]

    # PDF
    res = client.post("/api/exports/pdf", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"

    # PNG
    res = client.post("/api/exports/png", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "image/png"

    # Bundle ZIP
    res = client.post("/api/exports/bundle", json=payload)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/zip"

def test_api_preferences_crud_and_presets():
    client = TestClient(app)

    # 1. Get default preferences
    res = client.get("/api/preferences")
    assert res.status_code == 200
    data = res.json()
    assert data["schema_version"] == 2
    assert "theme" in data

    # 2. Save updated preferences
    save_payload = {
        "schema_version": 2,
        "theme": "light",
        "hard_constraints": {"max_daily_hours": 5.0, "day_pattern": "MORNING_ONLY"},
        "soft_preferences": {"faculty_rating_weight": 2.0},
        "selected_courses": ["CSE1001", "MAT1011"],
        "named_presets": {}
    }
    res = client.post("/api/preferences", json=save_payload)
    assert res.status_code == 200

    # 3. Add named preset
    preset_payload = {
        "name": "Custom Morning Fast",
        "hard_constraints": {"day_pattern": "MORNING_ONLY"},
        "soft_preferences": {"compactness_weight": 1.5}
    }
    res = client.post("/api/preferences/presets", json=preset_payload)
    assert res.status_code == 200

    # Verify preset persisted
    res = client.get("/api/preferences")
    assert res.status_code == 200
    named = res.json()["named_presets"]
    assert "Custom Morning Fast" in named

    # 4. Delete named preset
    res = client.delete("/api/preferences/presets/Custom Morning Fast")
    assert res.status_code == 200

    res = client.get("/api/preferences")
    assert "Custom Morning Fast" not in res.json()["named_presets"]

def test_api_database_system_endpoints():
    client = TestClient(app)

    # Integrity check
    res = client.get("/api/datasets/system/integrity")
    assert res.status_code == 200
    data = res.json()
    assert data["is_valid"] is True
    assert data["foreign_key_violations"] == 0

    # Export backup bundle
    res = client.get("/api/datasets/system/backup/bundle")
    assert res.status_code == 200
    bundle = res.json()
    assert bundle["schema_version"] == 1
    assert "favourites" in bundle
    assert "preferences" in bundle

    # Restore bundle
    res = client.post("/api/datasets/system/restore/bundle", json=bundle)
    assert res.status_code == 200
    assert res.json()["status"] == "restored"
